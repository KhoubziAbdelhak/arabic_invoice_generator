import os
from docx import Document

# * Directory paths
BASE_DIR = "/var/home/abdelhak/programming/pfe/arabic_invoice_generator/"
TEMPLATE_DIR = os.path.join(BASE_DIR, 'data', 'templates')
OUTPUT_DIR = os.path.join(BASE_DIR, 'output')
DOCX_DIR = os.path.join(OUTPUT_DIR, 'docx')
PDF_DIR = os.path.join(OUTPUT_DIR, 'pdf')
IMAGES_DIR = os.path.join(OUTPUT_DIR, 'images')
ANNOTATIONS_DIR = os.path.join(OUTPUT_DIR, 'annotations')
AUGMENTED_IMAGES_DIR = os.path.join(OUTPUT_DIR, 'augmented_images')
AUGMENTED_ANNOTATIONS_DIR = os.path.join(OUTPUT_DIR, 'augmented_annotations')

# * Create directories if they don't exist
for directory in [TEMPLATE_DIR, OUTPUT_DIR, DOCX_DIR, PDF_DIR, IMAGES_DIR, ANNOTATIONS_DIR]:
    os.makedirs(directory, exist_ok=True)

# * Files paths
PRODUCTS_CSV_PATH = os.path.join(BASE_DIR, 'data', 'products.csv')


# * Invoice configuration
VAT_RATE = 0.15
NUM_PRODUCTS_PER_INVOICE = 5
NUM_INVOICES_TO_GENERATE = 20
NUM_AUGMENTED_IMAGES = NUM_INVOICES_TO_GENERATE


# * Add template validation
def validate_template(template_path):
    try:
        doc = Document(template_path)
        required_placeholders = {
            '{{invoice_ref}}', '{{seller_name}}', '{{seller_address}}',
            '{{seller_vat_number}}', '{{issue_datetime}}', '{{products_table}}',
            '{{subtotal}}', '{{tax}}', '{{total}}', '{{email}}',
            '{{website}}', '{{phone_number}}', '{{fax_number}}'
        }

        found_placeholders = set()

        # Check paragraphs
        for paragraph in doc.paragraphs:
            for placeholder in required_placeholders:
                if placeholder in paragraph.text:
                    found_placeholders.add(placeholder)

        # Check tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for placeholder in required_placeholders:
                            if placeholder in paragraph.text:
                                found_placeholders.add(placeholder)

        missing = required_placeholders - found_placeholders
        if missing:
            print(f"Warning: Missing placeholders in template: {missing}")
            return False
        return True
    except Exception as e:
        print(f"Template validation failed: {e}")
        return False

# ! Validate templates on startup
for template_file in os.listdir(TEMPLATE_DIR):
    template_path = os.path.join(TEMPLATE_DIR, template_file)
    if validate_template(template_path):
        print(f"Template '{template_file}' is valid.")
    else:
        print(f"Warning: Template '{template_file}' is invalid.")
